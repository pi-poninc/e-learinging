#!/usr/bin/env python3
"""第二の脳 実演デモ（操作イメージ版）。
画面録画の代わりに、Obsidianを模したモック画面を各ステップで描く。8スライド。
台本(script_台本_cc.md)の行数=8 に一致させること。ワイプなしで、スライド＋ナレ＋テロップとして合成する。"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

JP = "Hiragino Kaku Gothic Pro"
MONO = "Menlo"
# Obsidian風ダークテーマ
BG      = RGBColor(0x1E, 0x1E, 0x1E)
SIDE    = RGBColor(0x26, 0x26, 0x26)
TITLEBAR= RGBColor(0x32, 0x32, 0x32)
PANE    = RGBColor(0x1E, 0x1E, 0x1E)
FG      = RGBColor(0xDC, 0xDD, 0xDE)
MUTED   = RGBColor(0x8A, 0x8A, 0x8A)
ACCENT  = RGBColor(0x8B, 0x7C, 0xF6)   # obsidian purple
ACCENTBG= RGBColor(0x2E, 0x2A, 0x47)
LINK    = RGBColor(0xA5, 0x8C, 0xFF)
GREEN   = RGBColor(0x6B, 0xD9, 0x68)
CARD    = RGBColor(0x2A, 0x2A, 0x2A)
EMU_W, EMU_H = Inches(13.333), Inches(7.5)
FOLDERS = ["00_Inbox", "10_Daily", "20_Projects", "30_Areas",
           "40_Resources", "50_PermanentNotes", "90_Archive"]


def _font(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)


def rect(s, x, y, w, h, fill, line=None, rounded=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    if rounded:
        try: shp.adjustments[0] = 0.08
        except Exception: pass
    return shp


def oval(s, x, y, d, fill):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill; shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def txt(s, x, y, w, h, runs, size=14, color=FG, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, font=JP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, color, bold)]
    p = tf.paragraphs[0]; p.alignment = align
    for seg in runs:
        t, c, b = (seg + (color, bold))[:3] if isinstance(seg, tuple) else (seg, color, bold)
        r = p.add_run(); r.text = t; f = r.font
        f.size = Pt(size); f.bold = b; f.color.rgb = c
        _font(r, font)
    return tb


def line(s, x1, y1, x2, y2, color, w=1.0):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color; cn.line.width = Pt(w); cn.shadow.inherit = False
    return cn


def window(s, step_label=None):
    """ウィンドウ枠＋タイトルバー＋サイドバーのガワを描き、本文領域の左上x,yを返す。"""
    rect(s, 0, 0, 13.333, 7.5, BG)
    rect(s, 0, 0, 13.333, 0.5, TITLEBAR)           # title bar
    for i, c in enumerate([RGBColor(0xFF,0x5F,0x56), RGBColor(0xFF,0xBD,0x2E), RGBColor(0x27,0xC9,0x3F)]):
        oval(s, 0.28 + i*0.32, 0.17, 0.17, c)
    txt(s, 1.5, 0.05, 6, 0.4, "Obsidian — 第二の脳", 13, MUTED, anchor=MSO_ANCHOR.MIDDLE)
    if step_label:                                  # タイトルバー右（カードに隠れない位置）
        b = rect(s, 10.95, 0.08, 2.15, 0.34, ACCENT, rounded=True)
        b.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = b.text_frame.paragraphs[0].add_run(); r.text = step_label
        r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); _font(r, JP)
    rect(s, 0, 0.5, 3.4, 7.0, SIDE)                 # sidebar
    txt(s, 0.3, 0.62, 3.0, 0.3, "VAULT ／ 第二の脳", 11, MUTED, bold=True)
    return 3.75, 1.0


def sidebar_folders(s, active=None):
    y = 1.08
    for i, f in enumerate(FOLDERS):
        if i == active:
            rect(s, 0.18, y-0.03, 3.05, 0.42, ACCENTBG, rounded=True)
        col = FG if i == active else MUTED
        txt(s, 0.42, y, 2.9, 0.4, [("", col, False), (f, col, i == active)], 13, col,
            bold=(i == active), anchor=MSO_ANCHOR.MIDDLE)
        y += 0.5


def notepane(s, title, body):
    """本文ペイン。body = [(kind, text)] kind: h1/h2/body/bullet/check/link/muted"""
    rect(s, 3.55, 0.62, 9.6, 0.5, CARD, rounded=True)       # tab
    txt(s, 3.75, 0.62, 9.2, 0.5, title, 13, FG, anchor=MSO_ANCHOR.MIDDLE)
    y = 1.35
    for kind, text in body:
        if kind == "h1":
            txt(s, 3.75, y, 9.2, 0.6, text, 30, FG, bold=True); y += 0.75
        elif kind == "h2":
            txt(s, 3.75, y, 9.2, 0.45, text, 18, ACCENT, bold=True); y += 0.55
        elif kind == "bullet":
            txt(s, 3.95, y, 9.0, 0.42, "・" + text, 16, FG); y += 0.5
        elif kind == "check":
            txt(s, 3.95, y, 9.0, 0.42, [("☐  ", MUTED, False), (text, FG, False)], 16, FG); y += 0.5
        elif kind == "link":
            txt(s, 3.95, y, 9.0, 0.42, [("[[", MUTED, False), (text, LINK, True), ("]]", MUTED, False)], 16, LINK); y += 0.5
        elif kind == "muted":
            txt(s, 3.75, y, 9.2, 0.42, text, 13, MUTED); y += 0.46
        else:
            txt(s, 3.75, y, 9.2, 0.42, text, 16, FG); y += 0.5
    return y


def build():
    prs = Presentation(); prs.slide_width = EMU_W; prs.slide_height = EMU_H
    blank = prs.slide_layouts[6]

    def new():
        return prs.slides.add_slide(blank)

    # 1) イントロ：Vault全景 + HOME
    s = new(); window(s); sidebar_folders(s, active=None)
    notepane(s, "HOME.md", [
        ("h1", "Second Brain — Home"),
        ("muted", "読むのは Obsidian、動かすのは AI（Claude Code）"),
        ("h2", "今日のクイックアクション"),
        ("bullet", "ためる → まわす → つなげる → 生かす"),
        ("muted", "左のフォルダは PARA：Inbox / Daily / Projects / Areas / Resources / PermanentNotes / Archive"),
    ])

    # 2) STEP1 Inbox にメモ
    s = new(); window(s, "STEP 1 ためる"); sidebar_folders(s, active=0)
    notepane(s, "00_Inbox / 2026-10-10 思いつき.md", [
        ("muted", "思いついたら、まずここに一行で放り込む（整理はしない）"),
        ("body", "・確認事項は「型」で出すと速い気がする"),
        ("body", "・システム名／確認事項／担当／期限／参考、の5項目"),
        ("muted", "— あとで Permanent に育てる —"),
    ])

    # 3) STEP2 Daily + Top3
    s = new(); window(s, "STEP 2 まわす"); sidebar_folders(s, active=1)
    notepane(s, "10_Daily / 2026-10-10.md", [
        ("h1", "2026-10-10"),
        ("h2", "今日やること"),
        ("check", "確認事項テンプレを1枚つくる  ← Top1"),
        ("check", "定例の論点を3つ出す  ← Top2"),
        ("check", "Permanent を1枚書く  ← Top3"),
    ])

    # 4) STEP3 Permanent Note
    s = new(); window(s, "STEP 3 つなげる"); sidebar_folders(s, active=5)
    notepane(s, "50_PermanentNotes / 確認は型で出す.md", [
        ("h1", "確認は型で出す"),
        ("h2", "一言でいうと"),
        ("body", "確認は、システム名・内容・担当・期限・参考の型に揃える"),
        ("h2", "自分の言葉で"),
        ("body", "相手の回答コストが下がり、未回答の管理もできる"),
        ("h2", "具体例 ／ 関連ノート"),
    ])

    # 5) STEP4 リンク
    s = new(); window(s, "STEP 4 結ぶ"); sidebar_folders(s, active=5)
    y = notepane(s, "50_PermanentNotes / 確認は型で出す.md", [
        ("h2", "関連ノート"),
        ("link", "確認事項マネジメント"),
        ("link", "決定事項シート"),
        ("muted", "[[ を打つと候補が出る。リンクをたどると関連知識へ飛べる"),
    ])
    rect(s, 3.75, y+0.1, 9.2, 1.3, CARD, rounded=True)
    txt(s, 3.95, y+0.25, 8.8, 0.4, "バックリンク（このノートを参照しているノート）", 12, MUTED, bold=True)
    txt(s, 3.95, y+0.7, 8.8, 0.4, [("・", MUTED, False), ("PM編：複数社連携の確認事項マネジメント", LINK, False)], 14, LINK)

    # 6) STEP5 グラフビュー
    s = new(); window(s, "STEP 5 俯瞰"); sidebar_folders(s, active=None)
    rect(s, 3.55, 0.62, 9.6, 6.6, RGBColor(0x16,0x16,0x16), rounded=True)
    txt(s, 3.75, 0.75, 9.2, 0.4, "グラフビュー ― 知識のつながり", 13, MUTED, bold=True)
    nodes = [(6.0,2.2,"確認は型で出す",True),(4.5,4.0,"確認事項\nマネジメント",False),
             (7.8,3.4,"決定事項\nシート",False),(9.8,2.3,"要件整理",False),
             (8.6,5.2,"中間納品",False),(5.3,5.6,"速いQA",False),(10.6,4.4,"PM編",False)]
    cx, cy = 6.0+0.55, 2.2+0.35
    for (x,y,_,_) in nodes[1:]:
        line(s, cx, cy, x+0.4, y+0.3, RGBColor(0x44,0x3F,0x66), 1.2)
    line(s, 4.5+0.4,4.0+0.3, 7.8+0.4,3.4+0.3, RGBColor(0x3A,0x36,0x55), 1.0)
    for (x,y,lbl,center) in nodes:
        d = 0.9 if center else 0.66
        oval(s, x, y, d, ACCENT if center else RGBColor(0x5B,0x53,0x8A))
        txt(s, x-0.7, y+d+0.02, d+1.4, 0.5, lbl, 11 if center else 10, FG if center else MUTED,
            bold=center, align=PP_ALIGN.CENTER)

    # 7) STEP6 AIに聞く（ターミナル）
    s = new(); window(s, "STEP 6 生かす")
    rect(s, 0, 0.5, 13.333, 7.0, RGBColor(0x14,0x14,0x14))
    txt(s, 0.6, 0.75, 12, 0.4, "● ● ●  claude — 第二の脳", 12, MUTED, font=MONO)
    tl = [
        ("$ claude", GREEN),
        ("> /ask-vault 今週の要点は？", FG),
        ("", FG),
        ("Vault を横断して回答します：", MUTED),
        ("  1. ある案件：中間納品は未確定事項一覧を添えて完了", FG),
        ("  2. 確認事項は「型」で出す運用に統一（PM編ノート）", FG),
        ("  3. 来週の論点：認証ロック要件を決裁定例で頭出し", FG),
        ("", FG),
        ("→ 関連ノート: [[確認は型で出す]] [[決定事項シート]]", LINK),
    ]
    y = 1.35
    for t, c in tl:
        txt(s, 0.7, y, 12.2, 0.38, t, 15, c, font=MONO); y += 0.5

    # 8) まとめ
    s = new(); window(s); sidebar_folders(s, active=None)
    notepane(s, "ダッシュボード.md", [
        ("h1", "第二の脳の1日"),
        ("h2", "ためる → つなげる → 引き出す"),
        ("bullet", "ためる：思いつきを Inbox へ（手で一行）"),
        ("bullet", "つなげる：自分の言葉で Permanent＋[[リンク]]"),
        ("bullet", "引き出す：AI に一言（/ask-vault・横断検索）"),
        ("muted", "小さく始めれば、今日からあなたも動かせる"),
    ])

    out = Path(__file__).resolve().parent / "slides.pptx"
    prs.save(str(out))
    print(f"  saved {out}  ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
