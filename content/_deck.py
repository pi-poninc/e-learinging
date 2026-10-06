#!/usr/bin/env python3
"""スライドデッキ共通ヘルパー。各フォルダの gen_slides.py は SLIDES を定義して build_deck() を呼ぶだけ。

既存レッスン（01-fde-genba / 02-workshop）と同じ見た目（ダーク背景＋オレンジのアクセント、
16:9、Hiragino）に揃えてある。1スライド=1メッセージが基本。

SLIDES の各要素は dict:
  {"type": "title", "kicker":..., "title":..., "sub":..., "foot":...}
  {"kicker":..., "title":..., "bullets": [..], "note": "任意の補足1行"}
  {"kicker":..., "title":..., "rows": [("左","右"), ...], "note": ...}   # 2列テーブル
bullets は "・" を自動で付ける。項目が多いと自動で小さめの文字にする。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

JP_FONT = "Hiragino Kaku Gothic Pro"

BG     = RGBColor(0x14, 0x14, 0x14)
FG     = RGBColor(0xF5, 0xF3, 0xF0)
ACCENT = RGBColor(0xE8, 0x7A, 0x3E)
MUTED  = RGBColor(0x9A, 0x94, 0x8C)
CARD   = RGBColor(0x20, 0x1F, 0x1D)
EMU_W, EMU_H = Inches(13.333), Inches(7.5)


def _apply_font(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def _para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, first=False, space_after=8):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    _apply_font(r, JP_FONT)
    return p


def _box(s, x, y, w, h):
    tb = s.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    return tb.text_frame


def _bg(prs, s, color=BG):
    r = s.shapes.add_shape(1, 0, 0, EMU_W, EMU_H)
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    s.shapes._spTree.remove(r._element)
    s.shapes._spTree.insert(2, r._element)


def build_deck(slides, out_path):
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    blank = prs.slide_layouts[6]

    for sl in slides:
        s = prs.slides.add_slide(blank)
        _bg(prs, s)

        if sl.get("type") == "title":
            tf = _box(s, Inches(0.9), Inches(2.1), Inches(11.6), Inches(3.4))
            if sl.get("kicker"):
                _para(tf, sl["kicker"], 22, ACCENT, bold=True, first=True, space_after=16)
                _para(tf, sl["title"], 42, FG, bold=True, space_after=8)
            else:
                _para(tf, sl["title"], 42, FG, bold=True, first=True, space_after=8)
            if sl.get("sub"):
                _para(tf, sl["sub"], 25, FG, space_after=18)
            if sl.get("foot"):
                _para(tf, sl["foot"], 18, MUTED)
            continue

        # kicker
        if sl.get("kicker"):
            _para(_box(s, Inches(0.9), Inches(0.7), Inches(11.6), Inches(0.6)),
                  sl["kicker"], 20, ACCENT, bold=True, first=True)
        # title
        _para(_box(s, Inches(0.9), Inches(1.35), Inches(11.6), Inches(1.9)),
              sl["title"], 34, FG, bold=True, first=True)

        if sl.get("rows"):
            rows = sl["rows"]
            y = 3.15
            left_w, right_x = 4.3, 5.5
            size = 22 if len(rows) <= 6 else 19
            gap = 0.72 if len(rows) <= 6 else 0.6
            for i, (l, r) in enumerate(rows):
                yy = Inches(y + i * gap)
                _para(_box(s, Inches(1.0), yy, Inches(left_w), Inches(gap)),
                      l, size, ACCENT, bold=True, first=True)
                _para(_box(s, Inches(right_x), yy, Inches(13.0 - right_x - 0.6), Inches(gap)),
                      r, size, FG, first=True)
        elif sl.get("bullets"):
            items = sl["bullets"]
            size = 23 if len(items) <= 4 else (21 if len(items) <= 5 else 19)
            gap = 14 if len(items) <= 4 else (11 if len(items) <= 5 else 8)
            tf = _box(s, Inches(1.0), Inches(3.15), Inches(11.4), Inches(3.9))
            for i, it in enumerate(items):
                _para(tf, "・" + it, size, FG, first=(i == 0), space_after=gap)

        if sl.get("note"):
            _para(_box(s, Inches(1.0), Inches(6.85), Inches(11.4), Inches(0.5)),
                  sl["note"], 17, MUTED, first=True)

    prs.save(str(out_path))
    print(f"  saved {out_path}  ({len(slides)} slides)")
