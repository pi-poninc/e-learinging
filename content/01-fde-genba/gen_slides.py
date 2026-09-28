#!/usr/bin/env python3
"""memo.md から起こしたスライドを PPTX 生成する（1スライド=1メッセージ）。
台本(script_台本_cc.md)の行数=このスライド枚数 に必ず合わせること。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# LibreOfficeで確実に解決できる日本語フォント（fc-listの alias に一致）。
# python-pptxのfont.nameはLatin(<a:latin>)しか設定しないので、East Asian(<a:ea>)/
# 複合(<a:cs>)にも同じ書体を明示しないと日本語グリフが消える。
JP_FONT = "Hiragino Kaku Gothic Pro"

def _apply_font(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)

# ブランドカラー（オレンジ系アクセント）
BG      = RGBColor(0x14, 0x14, 0x14)
FG      = RGBColor(0xF5, 0xF3, 0xF0)
ACCENT  = RGBColor(0xE8, 0x7A, 0x3E)  # テラコッタ寄りオレンジ
MUTED   = RGBColor(0x9A, 0x94, 0x8C)

EMU_W, EMU_H = Inches(13.333), Inches(7.5)  # 16:9

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]

def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(1, 0, 0, EMU_W, EMU_H)  # rectangle bg
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2, r._element)
    return s

def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    return tf

def para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, first=False, space_after=8):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space_after)
    r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    _apply_font(r, JP_FONT)
    return p

def kicker(s, text):
    tf = box(s, Inches(0.9), Inches(0.7), Inches(6), Inches(0.6))
    para(tf, text, 20, ACCENT, bold=True, first=True)

def title(s, text, y=1.4, size=40):
    tf = box(s, Inches(0.9), Inches(y), Inches(11.5), Inches(1.6))
    para(tf, text, size, FG, bold=True, first=True)

def bullets(s, items, y=3.2, size=24, gap=14):
    tf = box(s, Inches(1.0), Inches(y), Inches(11.3), Inches(3.6))
    for i, it in enumerate(items):
        para(tf, "・" + it, size, FG, first=(i == 0), space_after=gap)

# --- Slide 1: タイトル ---
s = slide()
tf = box(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.8))
para(tf, "新入社員向け 研修", 22, ACCENT, bold=True, first=True, space_after=16)
para(tf, "FDEの基本：", 46, FG, bold=True, space_after=4)
para(tf, "コードより先に、現場を見る", 46, FG, bold=True, space_after=20)
para(tf, "株式会社pipon｜社内教育コンテンツ  2026.09", 18, MUTED)

# --- Slide 2: FDEとは ---
s = slide(); kicker(s, "そもそも"); title(s, "FDE（フォワードデプロイドエンジニア）とは？")
bullets(s, [
    "顧客の「現場」に入り込んで、実際の課題を自分の手で解くエンジニア",
    "仕様書を待つのではなく、業務そのものを情報源にする",
    "作ることが目的ではなく、成果（業務が良くなること）が目的",
])

# --- Slide 3: 今日の気づき（起きたこと） ---
s = slide(); kicker(s, "現場で起きたこと"); title(s, "危うく「依頼書どおり」作りかけた")
bullets(s, [
    "「この作業を自動化したい」— 依頼メモを見た瞬間に実装イメージが湧いた",
    "その場でコードを書き始めそうになった",
    "でも一度手を止めて、まず現場を見ることにした",
])

# --- Slide 4: やったこと ---
s = slide(); kicker(s, "やったこと"); title(s, "1時間、隣で作業を観察した")
bullets(s, [
    "担当者の作業を、実際に隣で見せてもらった",
    "「当たり前すぎて言わないこと」を含めて聞いた",
    "自分でも同じ作業を少しやってみた",
])

# --- Slide 5: わかったこと ---
s = slide(); kicker(s, "わかったこと"); title(s, "本当のボトルネックは別だった")
bullets(s, [
    "依頼書の作業は、全体のごく一部でしかなかった",
    "本当に時間を食っていたのは『別Excelとの二重入力』",
    "依頼どおり作っていたら、効果はほぼゼロだった",
])

# --- Slide 6: 学び ---
s = slide(RGBColor(0x1E, 0x17, 0x11))
kicker(s, "新人へ、これだけは")
title(s, "見てから作る。作ってから聞かない。", size=38)
bullets(s, [
    "実装スピードより、最初の観察の1時間が成果を決める",
    "わからない時は、手を動かす前に現場を見る・聞く・やってみる",
    "ボトルネックは、現場の人が言わない所に埋まっている",
], y=3.4)

import os
out = os.path.join(os.path.dirname(__file__), "slides.pptx")
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
