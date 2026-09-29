#!/usr/bin/env python3
"""memo.md から起こした「顧客ワークショップの回し方」スライドをPPTX生成（1スライド=1メッセージ）。
台本(script_台本_cc.md)の行数=スライド枚数(6)に一致させること。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

JP_FONT = "Hiragino Kaku Gothic Pro"

def _apply_font(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)

BG     = RGBColor(0x14, 0x14, 0x14)
FG     = RGBColor(0xF5, 0xF3, 0xF0)
ACCENT = RGBColor(0xE8, 0x7A, 0x3E)
MUTED  = RGBColor(0x9A, 0x94, 0x8C)
EMU_W, EMU_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]

def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(1, 0, 0, EMU_W, EMU_H)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2, r._element)
    return s

def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(x, y, w, h); tb.text_frame.word_wrap = True
    return tb.text_frame

def para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, first=False, space_after=8):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space_after)
    r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    _apply_font(r, JP_FONT)
    return p

def kicker(s, text):
    para(box(s, Inches(0.9), Inches(0.7), Inches(8), Inches(0.6)), text, 20, ACCENT, bold=True, first=True)

def title(s, text, y=1.4, size=38):
    para(box(s, Inches(0.9), Inches(y), Inches(11.6), Inches(1.8)), text, size, FG, bold=True, first=True)

def bullets(s, items, y=3.3, size=23, gap=14):
    tf = box(s, Inches(1.0), Inches(y), Inches(11.4), Inches(3.6))
    for i, it in enumerate(items):
        para(tf, "・" + it, size, FG, first=(i == 0), space_after=gap)

# 1: タイトル
s = slide()
tf = box(s, Inches(0.9), Inches(2.3), Inches(11.6), Inches(3.0))
para(tf, "新入社員向け 研修", 22, ACCENT, bold=True, first=True, space_after=16)
para(tf, "会議に出るだけで、終わらせない", 42, FG, bold=True, space_after=6)
para(tf, "― 顧客ワークショップを「翌日の意思決定」に変える", 26, FG, space_after=18)
para(tf, "株式会社pipon｜pipon learning  2026.09", 18, MUTED)

# 2: 今日の実例
s = slide(); kicker(s, "今日の実例")
title(s, "JARMeC×キャミック 遠隔読影ワークショップ（第10回）")
bullets(s, [
    "遠隔読影の運用・例外処理・レポート入力の設計を議論",
    "曖昧なままだった運用を、その場で「構造」に落とした",
    "決めたことは、その日のうちにチームへ共有した",
])

# 3: 原則①
s = slide(); kicker(s, "原則①")
title(s, "曖昧な運用は「レーン / 階層」で構造化する")
bullets(s, [
    "例外処理は3レーンに：緊急・急ぎ / 情報不足 / 想定外の所見",
    "レポート入力は3階層に：自動入力 / 選択・定型文 / 自由記載",
    "「分類（レーン）」や「段階（階層）」に分けると、一気に決まる",
])

# 4: 原則②
s = slide(); kicker(s, "原則②")
title(s, "決定・未決・次アクションを分けて、即共有する")
bullets(s, [
    "決まったこと / まだ決まっていないこと / 次にやること を分ける",
    "その日のうちに共有する（鮮度が価値）",
    "分けるだけで、翌日から「誰が何をやるか」が明確になる",
])

# 5: 原則③
s = slide(); kicker(s, "原則③")
title(s, "AIに整理させ、人は「要らない所を削る」")
bullets(s, [
    "生の議事メモをAIに渡して、まず構造化させる",
    "人は、要らない所を削る・判断する に集中する",
    "「完璧な議事録」より「次に動ける要点」を優先する",
])

# 6: 学び
s = slide(RGBColor(0x1E, 0x17, 0x11)); kicker(s, "新人へ、これだけは")
title(s, "会議の成果は「議事録」ではなく「次に動ける状態」", size=34)
bullets(s, [
    "出て終わりにしない。翌日の意思決定を作る",
    "曖昧な業務は、レーン・階層に分けて構造化する",
    "決定・未決・次アクションを分けて、その日に共有する",
], y=3.4)

import os
out = os.path.join(os.path.dirname(__file__), "slides.pptx")
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
